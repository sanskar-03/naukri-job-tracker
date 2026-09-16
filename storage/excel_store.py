from pathlib import Path
from typing import Iterable, Dict, List
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["job_id", "title", "company", "location", "experience", "skills", "posted_date", "job_url", "scraped_at"]
MASTER_SHEET = "Master History"
HEADER_FILL = "1F2937"


def clean_sheet_name(keyword: str, location: str) -> str:
    name = f"{keyword} - {location}".strip()
    name = re.sub(r'[\\/*?:\[\]]', '-', name)
    name = re.sub(r'\s+', ' ', name).strip() or 'Search'
    return name[:31]


def ensure_workbook(path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = MASTER_SHEET
    ws.append(HEADERS)
    style_header(ws)
    prepare_sheet(ws)
    wb.save(path)
    wb.close()


def style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal='center', vertical='center')


def prepare_sheet(ws) -> None:
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    widths = {'A':18,'B':40,'C':30,'D':34,'E':18,'F':60,'G':18,'H':80,'I':22}
    for col, width in widths.items(): ws.column_dimensions[col].width = width


def _headers(ws) -> dict:
    return {str(c.value).strip(): i for i,c in enumerate(ws[1], start=1) if c.value}


def existing_keys(ws):
    hm = _headers(ws)
    id_col, url_col = hm.get('job_id'), hm.get('job_url')
    ids, urls = set(), set()
    if not id_col: return ids, urls
    for row in ws.iter_rows(min_row=2, values_only=True):
        if id_col <= len(row) and row[id_col-1]: ids.add(str(row[id_col-1]).strip())
        if url_col and url_col <= len(row) and row[url_col-1]: urls.add(str(row[url_col-1]).strip())
    return ids, urls


def read_sheet_jobs(ws) -> List[dict]:
    headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, '') for v in row): continue
        item = {}
        for i,h in enumerate(headers):
            if h: item[h] = row[i] if i < len(row) else ''
        out.append(item)
    return out


def read_all_jobs(path: str | Path) -> List[dict]:
    ensure_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    items = []
    for name in wb.sheetnames:
        if name != MASTER_SHEET:
            items.extend(read_sheet_jobs(wb[name]))
    wb.close()
    return items


def read_search_jobs(path: str | Path, keyword: str, location: str) -> List[dict]:
    ensure_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    name = clean_sheet_name(keyword, location)
    if name not in wb.sheetnames:
        wb.close(); return []
    rows = read_sheet_jobs(wb[name])
    wb.close(); return rows


def append_new_jobs(path: str | Path, jobs: Iterable[dict], keyword: str, location: str) -> List[dict]:
    ensure_workbook(path)
    path = Path(path)
    wb = load_workbook(path)
    master = wb[MASTER_SHEET]
    name = clean_sheet_name(keyword, location)
    if name in wb.sheetnames:
        search = wb[name]
    else:
        search = wb.create_sheet(name)
        for c, h in enumerate(HEADERS, start=1):
            search.cell(1, c).value = h
    if search['A1'].value != HEADERS[0]:
        for c, h in enumerate(HEADERS, start=1):
            search.cell(1, c).value = h
    style_header(search); style_header(master)
    prepare_sheet(search); prepare_sheet(master)
    search_ids, search_urls = existing_keys(search)
    master_ids, master_urls = existing_keys(master)
    new_jobs = []
    for job in jobs:
        jid = str(job.get('job_id') or '').strip()
        url = str(job.get('job_url') or '').strip()
        if not jid and not url: continue
        if (jid and jid in search_ids) or (url and url in search_urls): continue
        row = [job.get(h, '') for h in HEADERS]
        search.append(row)
        if jid: search_ids.add(jid)
        if url: search_urls.add(url)
        new_jobs.append(job)
        if not ((jid and jid in master_ids) or (url and url in master_urls)):
            master.append(row)
            if jid: master_ids.add(jid)
            if url: master_urls.add(url)
    prepare_sheet(search); prepare_sheet(master)
    wb.save(path); wb.close()
    return new_jobs


def search_history(path: str | Path) -> list[dict]:
    ensure_workbook(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    result=[]
    for name in wb.sheetnames:
        if name == MASTER_SHEET: continue
        result.append({'search':name, 'count': max(0, wb[name].max_row-1)})
    wb.close(); return result
